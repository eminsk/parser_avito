import tempfile
from pathlib import Path

from openpyxl import load_workbook

from models import Item, PriceDetailed, Geo
from parser.export.excel import ExcelStorage as Excel


def test_excel_save_handles_none_fields():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "test_export.xlsx"
        excel = Excel(file_path=file_path)

        ad_none = Item(
            id=1001,
            title="Тестовый товар",
            images=None,
            priceDetailed=None,
            sortTimeStamp=None,
            urlPath=None,
            coords=None,
            phone=None,
            sellerId=None,
            location=None,
            description=None,
            isPromotion=False,
            total_views=None,
            today_views=None
        )

        excel.save([ad_none])

        wb = load_workbook(file_path)
        sheet = wb.active
        assert sheet.max_row == 2  # header + 1 data row
        row_values = [cell.value for cell in sheet[2]]
        assert row_values[0] == "Тестовый товар"
        assert row_values[1] == 0  # Default price
        assert row_values[2] in ("", None)  # Empty URL when urlPath is None


def test_excel_save_preserves_zero_views():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "views_export.xlsx"
        excel = Excel(file_path=file_path)

        ad_zero_views = Item(
            id=1002,
            title="Товар 0 просмотров",
            total_views=0,
            today_views=0,
        )
        ad_none_views = Item(
            id=1003,
            title="Товар None просмотров",
            total_views=None,
            today_views=None,
        )

        excel.save([ad_zero_views, ad_none_views])

        wb = load_workbook(file_path)
        sheet = wb.active
        assert sheet.max_row == 3

        # Row 2: zero views
        row2 = [cell.value for cell in sheet[2]]
        # total_views is index 11, today_views is index 12
        assert row2[11] == 0
        assert row2[12] == 0

        # Row 3: None views should be empty string or None
        row3 = [cell.value for cell in sheet[3]]
        assert row3[11] in ("", None)
        assert row3[12] in ("", None)


def test_excel_formula_injection_prevention():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "formula_export.xlsx"
        excel = Excel(file_path=file_path)

        ad = Item(
            id=1004,
            title="=SUM(1, 2)",
            description="+cmd|' /C calc'!A0",
            phone="@admin",
            sellerId="-999"
        )

        excel.save([ad])

        wb = load_workbook(file_path)
        sheet = wb.active
        row = [cell.value for cell in sheet[2]]
        assert row[0] == "'=SUM(1, 2)"
        assert row[3] == "'+cmd|' /C calc'!A0"
        assert row[5] == "'-999"
        assert row[13] == "'@admin"


def test_excel_coords_handling():
    with tempfile.TemporaryDirectory() as tmp_dir:
        file_path = Path(tmp_dir) / "coords_export.xlsx"
        excel = Excel(file_path=file_path)

        ad_valid = Item(id=1, coords={"lat": 55.75, "lng": 37.61, "address_user": "Москва"})
        ad_none_coords = Item(id=2, coords=None)
        ad_null_lat = Item(id=3, coords={"lat": None, "lng": None})

        excel.save([ad_valid, ad_none_coords, ad_null_lat])

        wb = load_workbook(file_path)
        sheet = wb.active

        # Check coords (col index 8: 0-indexed)
        # Headers: [Заголовок, Цена, URL, Описание, Дата, Продавец, Адрес, Адрес пользователя, Координаты]
        row1 = [cell.value for cell in sheet[2]]
        row2 = [cell.value for cell in sheet[3]]
        row3 = [cell.value for cell in sheet[4]]

        assert row1[7] == "Москва"
        assert row1[8] == "55.75;37.61"
        assert row2[8] in ("", None)
        assert row3[8] in ("", None)
