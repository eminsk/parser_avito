from pathlib import Path
from threading import Lock
from datetime import datetime

from openpyxl import Workbook, load_workbook
from loguru import logger
from tzlocal import get_localzone

from parser.export.base import ResultStorage
from models import Item

class ExcelStorage(ResultStorage):
    """
    Сохранение результатов парсинга в XLSX
    """
    name = "excel"
    headers = [
        "Название",
        "Цена",
        "URL",
        "Описание",
        "Дата публикации",
        "Продавец",
        "Адрес",
        "Адрес пользователя",
        "Координаты",
        "Изображения",
        "Поднято",
        "Просмотры (всего)",
        "Просмотры (сегодня)",
        "Телефон"
    ]

    def __init__(self, file_path: Path):
        self.file_path = file_path

        # создаём директорию
        self.file_path.parent.mkdir(parents=True, exist_ok=True)

        # lock для потокобезопасной записи
        self._lock = Lock()

        # создаём файл, если его нет
        if not self.file_path.exists():
            self._create_file()

    def _create_file(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(self.headers)
        workbook.save(self.file_path)

    @staticmethod
    def _get_ad_time(ad: Item):
        if not ad.sortTimeStamp:
            return ""
        try:
            return (
                datetime
                .fromtimestamp(ad.sortTimeStamp / 1000, tz=get_localzone())
                .replace(tzinfo=None)
            )
        except Exception:
            return ""

    @staticmethod
    def _get_item_coords(ad: Item) -> str:
        if ad.coords and isinstance(ad.coords, dict):
            lat = ad.coords.get("lat")
            lng = ad.coords.get("lng")
            if lat is not None and lng is not None:
                return f"{lat};{lng}"
        return ""

    @staticmethod
    def _get_item_address_user(ad: Item) -> str:
        if ad.coords and isinstance(ad.coords, dict) and "address_user" in ad.coords:
            return str(ad.coords["address_user"] or "")
        return ""

    @staticmethod
    def _get_largest_image_url(img) -> str:
        try:
            if not getattr(img, "root", None):
                return ""
            def _calc_dim(k: str) -> int:
                parts = k.split("x")
                if len(parts) == 2 and parts[0].isdigit() and parts[1].isdigit():
                    return int(parts[0]) * int(parts[1])
                return 0
            best_key = max(img.root.keys(), key=_calc_dim)
            return str(img.root[best_key])
        except Exception as err:
            logger.error(f"При определении лучшего изображения ошибка: {err}")
            return ""

    @staticmethod
    def excel_safe(value):
        # Formula Injection fix
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    def save(self, ads: list[Item]) -> None:
        if not ads:
            return

        with self._lock:
            workbook = load_workbook(self.file_path)
            sheet = workbook.active

            for ad in ads:
                images_urls = [
                    self._get_largest_image_url(img)
                    for img in (ad.images or [])
                ]

                price_value = (
                    ad.priceDetailed.value
                    if (ad.priceDetailed and getattr(ad.priceDetailed, "value", None) is not None)
                    else 0
                )

                url_path = ad.urlPath or ""
                if not url_path.startswith("/") and url_path:
                    url_path = f"/{url_path}"
                item_url = f"https://www.avito.ru{url_path}" if url_path else ""

                row = [
                    self.excel_safe(ad.title),
                    price_value,
                    self.excel_safe(item_url),
                    self.excel_safe(ad.description),
                    self._get_ad_time(ad),
                    self.excel_safe(ad.sellerId or ""),
                    self.excel_safe(ad.location.name if ad.location else ""),
                    self.excel_safe(self._get_item_address_user(ad)),
                    self.excel_safe(self._get_item_coords(ad)),
                    self.excel_safe(";".join(images_urls)),
                    "Да" if ad.isPromotion else "Нет",
                    ad.total_views if ad.total_views is not None else "",
                    ad.today_views if ad.today_views is not None else "",
                    self.excel_safe(ad.phone or ""),
                ]

                sheet.append(row)

            workbook.save(self.file_path)



